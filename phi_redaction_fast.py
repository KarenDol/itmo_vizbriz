#!/usr/bin/env python3
"""
Fast PHI Redaction Script - Regex Only
High-performance batch processing for S3 Knowledge Base
"""

import os
import io
import re
import sys
import pathlib
import logging
import time
import gc
import psutil
from concurrent.futures import ThreadPoolExecutor, as_completed
import pdfplumber
from docx import Document
import boto3
from botocore.config import Config

# Set up logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('phi_redaction_fast.log'),
        logging.StreamHandler(sys.stdout)
    ]
)
logger = logging.getLogger(__name__)

# ---------- Configuration ----------
S3_KNOWLEDGE_BUCKET = os.environ.get('S3_BUCKET_NAME', 'vizbrizknowledgebase')
S3_OUTPUT_BUCKET = "vizbrizknowledgebase"
S3_OUTPUT_PREFIX = "redacted/"

# Allowed file types for processing
SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}

# ---------- Performance Configuration ----------
MAX_MEMORY_PERCENT = 80      # Stop if memory > 80%
MAX_FILE_SIZE_MB = 100       # Skip files > 100MB
MAX_TEXT_LENGTH = 2000000    # 2M character limit
THREAD_POOL_SIZE = 4         # Process 4 files simultaneously
BATCH_SIZE = 20              # Process 20 files per batch
BATCH_DELAY_SECONDS = 1      # 1 second delay between batches

# Terms that should never be redacted (medical/dental terminology)
ALLOWLIST = {
    "OSA", "CPAP", "MAD", "AHI", "RDI", "ODI", "SpO2", "PSG", "HSAT",
    "apnea", "hypopnea", "snoring", "sleep", "breathing", "airway",
    "mandibular", "advancement", "appliance", "titration", "compliance",
    "AASM", "AADSM", "FDA", "FDA-approved", "FDA cleared",
    "TMJ", "TMD", "bruxism", "clenching", "grinding", "occlusion"
}

def check_memory_usage():
    """Check current memory usage"""
    try:
        memory = psutil.virtual_memory()
        return memory.percent
    except Exception:
        return 0

def cleanup_memory():
    """Force garbage collection"""
    gc.collect()

def truncate_text_if_needed(text: str) -> str:
    """Truncate text if too long"""
    if len(text) > MAX_TEXT_LENGTH:
        logger.warning(f"Text truncated from {len(text)} to {MAX_TEXT_LENGTH} characters")
        return text[:MAX_TEXT_LENGTH] + "\n\n[TEXT TRUNCATED DUE TO SIZE LIMITS]"
    return text

def fast_redact_text(text: str) -> str:
    """Super-fast PHI redaction using only regex patterns"""
    try:
        # Protect allowlisted terms
        placeholders = {}
        for term in ALLOWLIST:
            key = f"__ALLOW_{hash(term) % 1000000}__"
            placeholders[key] = term
            text = re.sub(rf"\b{re.escape(term)}\b", key, text, flags=re.IGNORECASE)
        
        # Preserve age references
        age_placeholders = {}
        age_patterns = [
            # Common age formats
            r'\b(\d{1,3})\s*(?:yo|y\.o\.|year[s]?\s*old|yr[s]?\s*old)\b',  # "62yo", "62 years old"
            r'\b(?:age|aged?)\s*:?\s*(\d{1,3})\b',  # "age: 62", "aged 62"
            r'\((\d{1,3})yo\b',  # "(62yo"
            # Additional age patterns for better coverage
            r'\b(\d{1,3})\s*years?\s*of\s*age\b',  # "62 years of age"
            r'\b(\d{1,3})\s*y\.o\.\b',  # "62 y.o."
            r'\b(\d{1,3})\s*yrs?\s*old\b',  # "62 yrs old"
            r'\b(\d{1,3})\s*years?\b(?=\s*(?:old|of\s*age|male|female|patient))',  # "62 years" followed by context
            r'\b(?:patient|pt)\.?\s*(\d{1,3})\b',  # "patient 62" or "pt. 62"
            r'\b(\d{1,3})\s*(?:year|yr)\.?\s*old\b',  # "62 year old" or "62 yr. old"
            r'\b(\d{1,3})\s*yo\b',  # "62 yo" (standalone)
            r'\b(\d{1,3})\s*years?\s*old\b',  # "62 years old"
            r'\b(?:age|aged)\s*(\d{1,3})\b',  # "age 62" or "aged 62"
            r'\b(\d{1,3})\s*(?:year|yr)s?\s*of\s*age\b',  # "62 years of age"
        ]
        
        age_counter = 0
        for pattern in age_patterns:
            matches = re.finditer(pattern, text, flags=re.IGNORECASE)
            for match in matches:
                age_value = match.group(1)
                if 0 <= int(age_value) <= 120:
                    age_counter += 1
                    placeholder = f"__AGE_PRESERVE_{age_counter}__"
                    age_placeholders[placeholder] = match.group(0)
                    text = text.replace(match.group(0), placeholder, 1)
        
        # Apply fast PHI redaction patterns
        redacted_text = text
        
        # Patient names - more precise patterns
        # Pattern 1: Specific known patient name (KRISMAN, DENNIS)
        redacted_text = re.sub(r'\bKRISMAN,\s+DENNIS\b', '[PATIENT_NAME]', redacted_text)
        # Pattern 2: General LASTNAME, Firstname format (avoiding common words)
        redacted_text = re.sub(r'\b[A-Z]{2,20},\s+[A-Z][a-z]{2,20}(?:\s+[A-Z]\.?)?\b(?!\s*(?:ASSOC|SLEEP|PREMIER|MEDICAL|CONFIDENTIAL|GROUP))', '[PATIENT_NAME]', redacted_text)
        # Pattern 3: Names after specific labels (more precise)
        redacted_text = re.sub(r'\b(?:Name:|Patient Name:|RE:)\s*([A-Z][a-z]+\s+[A-Z][a-z]+(?:\s+[A-Z][a-z]+)?)\b', lambda m: m.group(0).replace(m.group(1), '[PATIENT_NAME]'), redacted_text)
        
        # Patient ID numbers - enhanced patterns
        redacted_text = re.sub(r'\b(?:id|ID|patient\s*(?:id|ID))\s*#?\s*:?\s*\d{3,}\b', '[PATIENT_ID]', redacted_text, flags=re.IGNORECASE)
        
        # Provider names - enhanced patterns
        # Pattern 1: Names with credentials
        redacted_text = re.sub(r'\b[A-Z][A-Z\s]+,?\s*(?:DDS|MD|NP|PA)\b', '[PROVIDER_NAME]', redacted_text)
        # Pattern 2: Dr. prefix
        redacted_text = re.sub(r'\bDr\.?\s+[A-Z][a-z]+\s+[A-Z][a-z]+(?:\s+[A-Z][a-z]+)?\b', '[PROVIDER_NAME]', redacted_text)
        # Pattern 3: Standalone professional names (like "Oleg K")
        redacted_text = re.sub(r'\b[A-Z][a-z]{2,15}\s+[A-Z](?:\.|[a-z]{1,15})?\b(?=\s*$|\s*\n|\s*[0-9])', '[PROVIDER_NAME]', redacted_text)
        # Pattern 4: Names followed by credentials on next line
        redacted_text = re.sub(r'\b([A-Z][a-z]+\s+[A-Z][a-z]+)\s*\n.*?(?:DDS|MD)', '[PROVIDER_NAME]', redacted_text, flags=re.MULTILINE)
        
        # Addresses - enhanced patterns
        # Full addresses with street, city, state, zip
        redacted_text = re.sub(r'\b\d{1,5}\s+[A-Za-z0-9.\'\-\s]+\s+(?:Street|St|Road|Rd|Ave|Avenue|Blvd|Boulevard|Dr|Drive|Lane|Ln|Court|Ct|Way|Place|Pl)(?:\s+(?:Apt|Suite|Unit|#)\s*[A-Za-z0-9\-]+)?(?:,?\s*[A-Za-z\s]+,?\s*[A-Z]{2}\s+\d{5}(?:-\d{4})?)?', '[ADDRESS]', redacted_text, flags=re.IGNORECASE)
        
        # City, State ZIP patterns (like "REDMOND, WA 98052")
        redacted_text = re.sub(r'\b[A-Z][A-Za-z\s]+,\s*[A-Z]{2}\s+\d{5}(?:-\d{4})?\b', '[CITY_STATE_ZIP]', redacted_text)
        
        # Standalone city names in address context
        redacted_text = re.sub(r'\b(?:City|Address).*?([A-Z][A-Za-z]+(?:\s+[A-Z][A-Za-z]+)*),?\s*[A-Z]{2}\b', lambda m: m.group(0).replace(m.group(1), '[CITY]'), redacted_text)
        
        # ZIP codes (standalone)
        redacted_text = re.sub(r'\b\d{5}(?:-\d{4})?\b', '[ZIP_CODE]', redacted_text)
        
        # Phone numbers
        redacted_text = re.sub(r'\b(?:\+?1[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}\b', '[PHONE_NUMBER]', redacted_text)
        
        # Email addresses
        redacted_text = re.sub(r'\b[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[A-Za-z]{2,}\b', '[EMAIL_ADDRESS]', redacted_text)
        
        # Insurance numbers
        redacted_text = re.sub(r'\b(?:Policy|Group|Member|Subscriber|Insurance)\s*(?:ID|No\.?|Number|#)\s*[:#]?\s*[A-Z0-9\-]{6,}\b', '[INSURANCE_POLICY]', redacted_text, flags=re.IGNORECASE)
        redacted_text = re.sub(r'\b[A-Z]{2,4}\d{6,12}\b', '[INSURANCE_NUMBER]', redacted_text)
        
        # SSN
        redacted_text = re.sub(r'\b\d{3}-?\d{2}-?\d{4}\b', '[SSN]', redacted_text)
        
        # Medical record numbers
        redacted_text = re.sub(r'\b(MRN|Med(?:ical)?\s*Record|Patient\s*ID)\s*#?:?\s*[A-Z0-9\-]{6,}\b', '[MEDICAL_RECORD]', redacted_text, flags=re.IGNORECASE)
        
        # Dates (preserve ages but redact DOB)
        redacted_text = re.sub(r'\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}[/-]\d{1,2}[/-]\d{1,2}|(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\s+\d{1,2},?\s+\d{4})\b', '[DATE]', redacted_text, flags=re.IGNORECASE)
        
        # Restore allowlisted terms
        for key, term in placeholders.items():
            redacted_text = redacted_text.replace(key, term)
        
        # Restore age references
        for placeholder, age_text in age_placeholders.items():
            redacted_text = redacted_text.replace(placeholder, age_text)
        
        return redacted_text
        
    except Exception as e:
        logger.error(f"Error in fast redaction: {e}")
        return text

def extract_text_from_file(file_content: bytes, file_extension: str) -> str:
    """Extract text from different file types - optimized"""
    try:
        if file_extension in {".txt", ".md"}:
            return file_content.decode("utf-8", errors="ignore")
        
        elif file_extension == ".docx":
            import tempfile
            with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp_file:
                tmp_file.write(file_content)
                tmp_file.flush()
                
                doc = Document(tmp_file.name)
                text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
                
                os.unlink(tmp_file.name)
                return text
        
        elif file_extension == ".pdf":
            import tempfile
            with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp_file:
                tmp_file.write(file_content)
                tmp_file.flush()
                
                text_chunks = []
                with pdfplumber.open(tmp_file.name) as pdf:
                    for page in pdf.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text_chunks.append(page_text)
                
                os.unlink(tmp_file.name)
                return "\n\n".join(text_chunks)
        
        else:
            return ""
            
    except Exception as e:
        logger.error(f"Error extracting text from {file_extension} file: {e}")
        return ""

def process_single_file(s3_client, bucket: str, key: str) -> tuple:
    """Process a single file - optimized for speed"""
    try:
        # Check file size
        response = s3_client.head_object(Bucket=bucket, Key=key)
        file_size_mb = response['ContentLength'] / (1024 * 1024)
        
        if file_size_mb > MAX_FILE_SIZE_MB:
            return (key, "skipped", f"File too large ({file_size_mb:.1f}MB)")
        
        # Get file extension
        file_extension = pathlib.Path(key).suffix.lower()
        if file_extension not in SUPPORTED_EXTENSIONS:
            return (key, "skipped", "Unsupported file type")
        
        # Create output key
        safe_key = key.replace('/', '_').replace('\\', '_')
        base_name = pathlib.Path(safe_key).stem
        output_key = f"{S3_OUTPUT_PREFIX}{base_name}.md"
        
        # Download and extract text
        response = s3_client.get_object(Bucket=bucket, Key=key)
        file_content = response['Body'].read()
        
        original_text = extract_text_from_file(file_content, file_extension)
        if not original_text or not original_text.strip():
            return (key, "skipped", "No text content")
        
        # Truncate if needed
        original_text = truncate_text_if_needed(original_text)
        
        # Fast redaction
        redacted_text = fast_redact_text(original_text)
        
        # Upload redacted content
        s3_client.put_object(
            Bucket=S3_OUTPUT_BUCKET,
            Key=output_key,
            Body=redacted_text.encode('utf-8'),
            ContentType='text/markdown',
            Metadata={
                'original-file': key,
                'phi-redacted': 'true' if original_text != redacted_text else 'false'
            }
        )
        
        # Clean up memory
        del original_text, redacted_text, file_content
        
        return (key, "success", output_key)
        
    except Exception as e:
        return (key, "error", str(e))

def process_files_batch(s3_client, files_batch: list) -> list:
    """Process a batch of files using thread pool"""
    results = []
    
    with ThreadPoolExecutor(max_workers=THREAD_POOL_SIZE) as executor:
        # Submit all files in batch
        future_to_file = {
            executor.submit(process_single_file, s3_client, S3_KNOWLEDGE_BUCKET, key): key 
            for key in files_batch
        }
        
        # Collect results
        for future in as_completed(future_to_file):
            key = future_to_file[future]
            try:
                result = future.result()
                results.append(result)
            except Exception as e:
                results.append((key, "error", str(e)))
    
    return results

def main():
    """Main function - fast batch processing"""
    try:
        logger.info("Starting FAST PHI redaction script")
        start_time = time.time()
        
        # Initialize S3
        region = os.environ.get('AWS_REGION', 'us-west-2')
        s3_client = boto3.client('s3', region_name=region, config=Config(signature_version='s3v4'))
        logger.info("S3 client initialized")
        
        # Get all files to process
        logger.info("Scanning S3 bucket for files...")
        files_to_process = []
        paginator = s3_client.get_paginator('list_objects_v2')
        
        for page in paginator.paginate(Bucket=S3_KNOWLEDGE_BUCKET):
            if 'Contents' not in page:
                continue
                
            for obj in page['Contents']:
                key = obj['Key']
                
                # Skip redacted files and directories
                if key.startswith(S3_OUTPUT_PREFIX) or key.endswith('/'):
                    continue
                
                # Check if supported file type
                file_extension = pathlib.Path(key).suffix.lower()
                if file_extension in SUPPORTED_EXTENSIONS:
                    files_to_process.append(key)
        
        total_files = len(files_to_process)
        logger.info(f"Found {total_files} files to process")
        
        # Process files in batches
        processed_count = 0
        skipped_count = 0
        error_count = 0
        
        for i in range(0, total_files, BATCH_SIZE):
            batch = files_to_process[i:i + BATCH_SIZE]
            batch_num = (i // BATCH_SIZE) + 1
            total_batches = (total_files + BATCH_SIZE - 1) // BATCH_SIZE
            
            logger.info(f"Processing batch {batch_num}/{total_batches} ({len(batch)} files)")
            
            # Check memory before processing batch
            memory_percent = check_memory_usage()
            if memory_percent > MAX_MEMORY_PERCENT:
                logger.error(f"Memory usage too high ({memory_percent:.1f}%), stopping")
                break
            
            # Process batch
            results = process_files_batch(s3_client, batch)
            
            # Count results
            for key, status, details in results:
                if status == "success":
                    processed_count += 1
                    logger.info(f"✓ {key} -> {details}")
                elif status == "skipped":
                    skipped_count += 1
                    logger.info(f"⚠ Skipped {key}: {details}")
                else:
                    error_count += 1
                    logger.error(f"✗ Error {key}: {details}")
            
            # Progress report
            total_processed = processed_count + skipped_count + error_count
            progress = (total_processed / total_files) * 100
            elapsed = time.time() - start_time
            rate = total_processed / elapsed if elapsed > 0 else 0
            
            logger.info(f"Progress: {total_processed}/{total_files} ({progress:.1f}%) - Rate: {rate:.1f} files/sec")
            
            # Cleanup and delay between batches
            cleanup_memory()
            if i + BATCH_SIZE < total_files:  # Don't delay after last batch
                time.sleep(BATCH_DELAY_SECONDS)
        
        # Final report
        elapsed = time.time() - start_time
        logger.info("=== FAST PHI REDACTION COMPLETE ===")
        logger.info(f"Total time: {elapsed:.1f} seconds ({elapsed/60:.1f} minutes)")
        logger.info(f"Files processed: {processed_count}")
        logger.info(f"Files skipped: {skipped_count}")
        logger.info(f"Files with errors: {error_count}")
        logger.info(f"Average rate: {(processed_count + skipped_count + error_count)/elapsed:.1f} files/sec")
        
        return 0
        
    except Exception as e:
        logger.error(f"Script failed: {e}")
        return 1

if __name__ == "__main__":
    sys.exit(main())
